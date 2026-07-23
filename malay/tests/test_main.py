import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook


MALAY_DIR = Path(__file__).resolve().parents[1]
SAMPLE_DIR = MALAY_DIR / "sample_files"
sys.path.insert(0, str(MALAY_DIR))

import main


SOURCE_WORKBOOK = (
    SAMPLE_DIR / "MANNAT-EP-474 FOR MANNAT TRANSLATION - ORiginal.xlsx"
)
CHARACTER_LIST = SAMPLE_DIR / "MANNAT - CHARACTER LIST.xlsx"
EPISODE_SYNOPSIS = SAMPLE_DIR / "MANNAT-EP-474 EPISODE SYNOPSIS.docx"
SERIES_SYNOPSIS = SAMPLE_DIR / "MANNAT - SERIES SYNOPSIS.docx"


class WorkbookParsingTests(unittest.TestCase):
    def test_sample_workbook_has_expected_rows_and_batches(self):
        workbook, plans = main.read_source_workbook(SOURCE_WORKBOOK)
        self.addCleanup(workbook.close)

        self.assertEqual(len(plans), 1)
        self.assertEqual(plans[0].sheet_name, "MANNAT-EP-474 V02")
        self.assertEqual(plans[0].header_row, 2)
        self.assertEqual(len(plans[0].records), 291)

        batches = main.split_records_into_batches(plans[0].records)
        self.assertEqual([len(batch.records) for batch in batches], [50, 50, 50, 50, 50, 41])
        self.assertEqual([row.row_id for row in batches[0].context_after], [51, 52, 53, 54])
        self.assertEqual([row.row_id for row in batches[1].context_before], [47, 48, 49, 50])

    def test_screenplay_rows_are_carried_into_following_dialogue(self):
        workbook, plans = main.read_source_workbook(SOURCE_WORKBOOK)
        self.addCleanup(workbook.close)

        first = plans[0].records[0]
        self.assertEqual(first.excel_row, 4)
        self.assertIn("OPENING SHOT", first.screenplay)
        self.assertEqual(first.english_subtitle, "Your private investigator,//Ronak... - Yes.")

    def test_supporting_documents_are_read(self):
        context = main.load_supporting_context(
            character_list_path=CHARACTER_LIST,
            episode_synopsis_path=EPISODE_SYNOPSIS,
            series_synopsis_path=SERIES_SYNOPSIS,
        )

        self.assertIn("MANNAT", context["character_list"])
        self.assertIn("charity event", context["episode_synopsis"])
        self.assertIn("Marine Drive", context["series_synopsis"])


class GeminiValidationTests(unittest.TestCase):
    def setUp(self):
        workbook, plans = main.read_source_workbook(SOURCE_WORKBOOK)
        self.addCleanup(workbook.close)
        self.records = plans[0].records[:2]

    def test_result_is_reordered_by_row_id(self):
        result = [
            {
                "row_id": self.records[1].row_id,
                "malay_translation": "Saya bertemu dengannya.",
            },
            {
                "row_id": self.records[0].row_id,
                "malay_translation": "Penyiasat peribadi awak,//Ronak... - Ya.",
            },
        ]
        translations = main.normalize_batch_result(result, 1, self.records)
        self.assertEqual(set(translations), {1, 2})

    def test_changed_line_break_markers_are_rejected(self):
        result = [
            {
                "row_id": self.records[0].row_id,
                "malay_translation": "Penanda ini telah dibuang.",
            },
            {
                "row_id": self.records[1].row_id,
                "malay_translation": "Saya bertemu dengannya.",
            },
        ]
        with self.assertRaisesRegex(ValueError, "line breaks"):
            main.normalize_batch_result(result, 1, self.records)


class EndToEndWorkbookTests(unittest.TestCase):
    def test_sample_output_preserves_rows_and_adds_malay_column(self):
        context = main.load_supporting_context(
            character_list_path=CHARACTER_LIST,
            episode_synopsis_path=EPISODE_SYNOPSIS,
            series_synopsis_path=SERIES_SYNOPSIS,
        )

        def fake_gemini_call(client, batch, supporting_context, **_kwargs):
            self.assertIs(supporting_context, context)
            return [
                {
                    "row_id": record.row_id,
                    "malay_translation": f"BM: {record.english_subtitle}",
                }
                for record in batch.records
            ]

        with tempfile.TemporaryDirectory() as temporary_directory:
            with patch("main.call_gemini_for_batch", side_effect=fake_gemini_call) as mocked:
                output_path = main.process_excel_file(
                    file_path=SOURCE_WORKBOOK,
                    supporting_context=context,
                    client=object(),
                    output_dir=temporary_directory,
                    sleep_fn=lambda _seconds: None,
                )

            self.assertEqual(mocked.call_count, 6)
            self.assertTrue(output_path.is_file())

            source = load_workbook(SOURCE_WORKBOOK)
            output = load_workbook(output_path)
            self.addCleanup(source.close)
            self.addCleanup(output.close)
            source_sheet = source.active
            output_sheet = output.active

            self.assertEqual(output_sheet.max_row, source_sheet.max_row)
            self.assertEqual(output_sheet.cell(2, 8).value, main.MALAY_HEADER)
            self.assertEqual(
                output_sheet.cell(4, 8).value,
                "BM: Your private investigator,//Ronak... - Yes.",
            )
            self.assertEqual(output_sheet.cell(314, 8).value, None)
            self.assertEqual(output_sheet.cell(4, 7).value, source_sheet.cell(4, 7).value)
            self.assertEqual(output_sheet.cell(4, 8).style_id, output_sheet.cell(4, 7).style_id)


if __name__ == "__main__":
    unittest.main()
